"""Add one slide to Final_Presentation_Speaker_Sentiment.pptx: a playable clip
shown moving through the pipeline, ending in the speaker verdict.

The slide is built the way figs/dataset_sample.png in the report is built, by
showing the data at each stage rather than describing it: a real waveform, the
real diarization strip, the real word timings, the real Bangla turn with the
judge's cited span highlighted, and the verdict as a mark on a stance scale.
Prose is kept to stage labels.

It is inserted after slide 3 (the technical pipeline framework), which is
where an audience has just seen the pipeline in the abstract. It is numbered
"03a" so that no other slide's printed number has to change: the numbers in
this deck are literal text, not a slide-number field, so renumbering would
mean editing every later slide.

No existing slide is read from, written to or reordered.

    python -X utf8 presentation_assets/build_clip_assets.py   # clip + envelope
    python -X utf8 presentation_assets/add_flow_slide.py      # this slide
"""
from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
# optional first argument: build into a staging copy instead of the live deck
# (PowerPoint locks the file while it is open)
DECK = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Final_Presentation_Speaker_Sentiment.pptx"
ASSETS = ROOT / "presentation_assets"
CLIP = ASSETS / "demo_clip_2107006_ep005.mp3"
POSTER = ASSETS / "audio_poster.png"
ENVELOPE = ASSETS / "clip_envelope.json"
INSERT_AFTER = 3
SLIDE_LABEL = "03a"

# ---- the deck's own design tokens, read off slide 3 -----------------------
ACCENT_DARK = RGBColor(0x1D, 0x4E, 0x5F)
ACCENT = RGBColor(0x2E, 0x8F, 0xA3)
CARD = RGBColor(0xF8, 0xF9, 0xF7)
RULE = RGBColor(0xC3, 0xCB, 0xC6)
INK = RGBColor(0x14, 0x18, 0x1B)
BODY = RGBColor(0x3E, 0x4A, 0x4F)
MUTED = RGBColor(0x6B, 0x78, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NEG = RGBColor(0xB5, 0x43, 0x2A)
POS = RGBColor(0x3C, 0x7D, 0x52)
NEU = RGBColor(0x6B, 0x78, 0x7D)
MIX = RGBColor(0xB5, 0x70, 0x1F)
HILITE = "FFE9A8"

MONO, SANS, BANGLA = "Consolas", "Segoe UI", "Nirmala UI"

# The waveform and the speaker strip share this x range, so both read against
# the same time axis and the handover lines up on each.
CLIP_START, CLIP_DUR = 2326.0, 14.0
HANDOVER = 2328.82
TRACK_X, TRACK_W = 2.30, 9.05
LEFT, CONTENT_W = 0.72, 11.89


def sec_to_x(t: float) -> float:
    return TRACK_X + TRACK_W * (t - CLIP_START) / CLIP_DUR


def textbox(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, spacing=1.0):
    """runs: (text, size, bold, colour, font[, highlight_hex])."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = "flow-text"
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    para.line_spacing = spacing
    for item in runs:
        text, size, bold, colour, font = item[:5]
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = colour
        run.font.name = font
        if len(item) > 5 and item[5]:
            # run-level highlight has no python-pptx API
            rpr = run.font._rPr
            hl = rpr.makeelement(qn("a:highlight"), {})
            hl.append(hl.makeelement(qn("a:srgbClr"), {"val": item[5]}))
            rpr.append(hl)
    return box


def rect(slide, x, y, w, h, fill, line=None, kind=MSO_SHAPE.RECTANGLE, line_w=17780):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Emu(line_w)
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def band(slide, y, h, label):
    """A stage band: full-width card with its label on the top-left."""
    rect(slide, LEFT, y, CONTENT_W, h, CARD, ACCENT_DARK)
    textbox(slide, LEFT + 0.16, y + 0.09, 5.0, 0.22, [(label, 11, True, MUTED, MONO)])


def pill(slide, x, y, w, h, text, colour, size=13):
    rect(slide, x, y, w, h, colour, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(slide, x, y + (h - 0.22) / 2, w, 0.22,
            [(text, size, True, WHITE, SANS)], align=PP_ALIGN.CENTER)


def build() -> None:
    prs = Presentation(str(DECK))
    before = len(prs.slides)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    env = json.loads(ENVELOPE.read_text(encoding="utf-8"))["bars"]

    # ------------------------------------------------------------- header
    textbox(slide, 11.61, 7.02, 1.0, 0.27,
            [(SLIDE_LABEL, 16, False, MUTED, MONO)], align=PP_ALIGN.RIGHT)
    textbox(slide, LEFT, 0.40, CONTENT_W, 0.30,
            [("LIVE DEMO / ONE CLIP THROUGH THE PIPELINE", 18, True, ACCENT, MONO)])
    textbox(slide, LEFT, 0.78, CONTENT_W, 0.46,
            [("Fourteen seconds, from sound to a speaker verdict", 32, True, INK, SANS)])

    # ============================================================= band 1
    y1, h1 = 1.42, 1.66
    band(slide, y1, h1, "1 / AUDIO, AND WHO IS SPEAKING")

    audio = slide.shapes.add_movie(
        str(CLIP), Inches(LEFT + 0.30), Inches(y1 + 0.46), Inches(0.72), Inches(0.72),
        poster_frame_image=str(POSTER), mime_type="audio/mpeg")
    audio.name = "demo-audio"
    textbox(slide, LEFT + 0.08, y1 + 1.24, 1.16, 0.20,
            [("click to play", 10, True, ACCENT, MONO)], align=PP_ALIGN.CENTER)

    # waveform, measured from the clip itself
    wave_mid, wave_max = y1 + 0.72, 0.38
    bar_w = TRACK_W / len(env) * 0.62
    for i, amp in enumerate(env):
        bh = max(0.020, amp * wave_max)
        bx = TRACK_X + TRACK_W * i / len(env)
        t = CLIP_START + CLIP_DUR * i / len(env)
        rect(slide, bx, wave_mid - bh / 2, bar_w, bh,
             ACCENT_DARK if t >= HANDOVER else MUTED)

    # diarization strip, same time axis
    strip_y, strip_h = y1 + 1.18, 0.21
    hx = sec_to_x(HANDOVER)
    rect(slide, TRACK_X, strip_y, hx - TRACK_X - 0.02, strip_h, MUTED)
    rect(slide, hx, strip_y, TRACK_X + TRACK_W - hx, strip_h, ACCENT_DARK)
    textbox(slide, TRACK_X + 0.05, strip_y + 0.025, 1.3, 0.18,
            [("speaker 0", 10, True, WHITE, MONO)])
    textbox(slide, hx + 0.07, strip_y + 0.025, 2.6, 0.18,
            [("speaker 3, from 2328.8 s", 10, True, WHITE, MONO)])
    # the handover, marked once and carried up through the waveform
    rect(slide, hx - 0.008, y1 + 0.34, 0.016, strip_y - (y1 + 0.34), MIX)
    textbox(slide, TRACK_X, strip_y + 0.23, 1.2, 0.18,
            [("2326.0 s", 10, False, MUTED, MONO)])
    textbox(slide, TRACK_X + TRACK_W - 1.2, strip_y + 0.23, 1.2, 0.18,
            [("2340.0 s", 10, False, MUTED, MONO)], align=PP_ALIGN.RIGHT)

    # ============================================================= band 2
    y2, h2 = 3.20, 1.00
    band(slide, y2, h2, "2 / WORDS, EACH WITH A TIMESTAMP")
    words = [("একটু", "2328.58"),
             ("খেয়াল", "2329.34"),
             ("করে", "2329.60"),
             ("দেখি", "2329.76"),
             ("যে", "2330.06"),
             ("আমরা", "2330.28"),
             ("এখানে", "2330.56"),
             ("অনেকবার", "2330.86")]
    cx = LEFT + 0.24
    for word, stamp in words:
        cw = 0.50 + 0.105 * len(word)
        rect(slide, cx, y2 + 0.34, cw, 0.48, WHITE, RULE, line_w=9525)
        textbox(slide, cx, y2 + 0.38, cw, 0.26,
                [(word, 15, False, INK, BANGLA)], align=PP_ALIGN.CENTER)
        textbox(slide, cx, y2 + 0.63, cw, 0.16,
                [(stamp, 9, False, MUTED, MONO)], align=PP_ALIGN.CENTER)
        cx += cw + 0.14
    textbox(slide, cx + 0.04, y2 + 0.46, 1.6, 0.26,
            [("… 33 in all", 13, False, MUTED, SANS)])

    # ============================================================= band 3
    y3, h3 = 4.32, 1.06
    band(slide, y3, h3, "3 / ONE SPEAKER TURN, AND WHAT IT EXPRESSES")
    textbox(slide, LEFT + 0.24, y3 + 0.40, 8.50, 0.56,
            [("… সেই লক্ষ্যে "
              "যেটা আমরা দেখি "
              "সেরকম কিন্তু ",
              17, False, INK, BANGLA),
             ("বড় কোন পরিবর্তন "
              "হয়নি", 17, True, INK, BANGLA, HILITE)],
            spacing=1.12)
    pill(slide, 9.55, y3 + 0.40, 1.34, 0.42, "negative", NEG, size=14)
    textbox(slide, 11.02, y3 + 0.49, 1.50, 0.24,
            [("one of his eight turns", 11, False, MUTED, SANS)])

    # ============================================================= band 4
    y4, h4 = 5.52, 1.32
    band(slide, y4, h4, "4 / THE SPEAKER, ACROSS THE WHOLE EPISODE")

    pill(slide, LEFT + 0.24, y4 + 0.46, 1.34, 0.50, "mixed", MIX, size=16)

    seq = [POS, POS, NEU, NEU, NEG, NEG, POS, NEU]
    tx = LEFT + 1.82
    textbox(slide, tx, y4 + 0.36, 2.8, 0.18,
            [("his eight turns, in order", 10, True, MUTED, MONO)])
    for colour in seq:
        rect(slide, tx, y4 + 0.58, 0.27, 0.32, colour, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
        tx += 0.33
    textbox(slide, LEFT + 1.82, y4 + 0.96, 3.4, 0.20,
            [("3 positive   3 neutral   2 negative", 11, False, BODY, SANS)])

    sx, sw, sy = 6.20, 3.90, y4 + 0.74
    textbox(slide, sx, y4 + 0.36, 3.2, 0.18,
            [("stance across the episode", 10, True, MUTED, MONO)])
    rect(slide, sx, sy, sw, 0.05, RULE)
    for frac, lab in ((0.0, "−1"), (0.5, "0"), (1.0, "+1")):
        rect(slide, sx + sw * frac - 0.008, sy - 0.05, 0.016, 0.15, RULE)
        textbox(slide, sx + sw * frac - 0.25, sy + 0.15, 0.50, 0.18,
                [(lab, 10, False, MUTED, MONO)], align=PP_ALIGN.CENTER)
    mark = sx + sw * (0.15 + 1) / 2
    rect(slide, mark - 0.055, sy - 0.12, 0.11, 0.29, MIX,
         kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    # to the right of the marker, clear of both the caption and the +1 tick
    textbox(slide, mark + 0.12, sy - 0.10, 0.80, 0.22,
            [("+0.15", 12, True, MIX, MONO)])

    textbox(slide, 10.42, y4 + 0.50, 2.05, 0.66,
            [("one turn says negative, the speaker overall does not",
              11, False, MUTED, SANS)], spacing=1.12)

    # ------------------------------------------ place, leaving others alone
    lst = prs.slides._sldIdLst
    entries = list(lst)
    lst.remove(entries[-1])
    lst.insert(INSERT_AFTER, entries[-1])

    prs.save(str(DECK))
    print(f"slides: {before} -> {len(Presentation(str(DECK)).slides)}; "
          f"new slide at position {INSERT_AFTER + 1}")


def patch_audio() -> None:
    """Turn the movie python-pptx wrote into a real audio object.

    python-pptx only exposes add_movie, which emits a videoFile reference and
    a p:video timing node. PowerPoint needs audioFile, an /audio relationship
    and a p:audio node, or the clip shows as a video placeholder.
    """
    tmp = DECK.with_suffix(".patching.pptx")
    shutil.copy2(DECK, tmp)
    src = zipfile.ZipFile(tmp)

    target = None
    for name in src.namelist():
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
            if "ppaction://media" in src.read(name).decode("utf-8"):
                target = name
    if target is None:
        raise SystemExit("could not find the slide holding the media object")
    rels = target.replace("slides/", "slides/_rels/") + ".rels"

    out = zipfile.ZipFile(DECK, "w", zipfile.ZIP_DEFLATED)
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename == target:
            text = data.decode("utf-8")
            text = text.replace("<a:videoFile ", "<a:audioFile ")
            text = text.replace("</a:videoFile>", "</a:audioFile>")
            text = text.replace("<p:video>", "<p:audio>").replace("</p:video>", "</p:audio>")
            data = text.encode("utf-8")
        elif item.filename == rels:
            data = data.decode("utf-8").replace(
                "officeDocument/2006/relationships/video",
                "officeDocument/2006/relationships/audio").encode("utf-8")
        out.writestr(item, data)
    out.close()
    src.close()
    tmp.unlink()
    print(f"patched {target}: videoFile -> audioFile, /video -> /audio, p:video -> p:audio")


if __name__ == "__main__":
    for path in (DECK, CLIP, POSTER, ENVELOPE):
        if not path.is_file():
            raise SystemExit(f"missing: {path}")
    build()
    patch_audio()
